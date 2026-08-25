from __future__ import annotations

import io
import os
from typing import Any
from deeptutor.services.custom.volunteer_table_dao import get_plan

TIER_LABELS = {"reach": "冲刺", "steady": "稳妥", "safe": "保底"}


def _slot_prob(s: dict) -> float:
    return s.get("group_prob", s.get("admission_prob", 0))


def _slot_majors_text(s: dict) -> str:
    majors = s.get("majors", [])
    if not majors:
        return ""
    parts = [f"{m.get('major_name','')}({m.get('tag','')})" for m in majors[:5]]
    return ", ".join(parts)


def _register_cjk_fonts() -> bool:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    CANDIDATES = [
        "/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf",
        "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
        "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
    ]
    for p in CANDIDATES:
        if os.path.exists(p):
            try:
                pdfmetrics.registerFont(TTFont("CJK-Regular", p))
                pdfmetrics.registerFont(TTFont("CJK-Bold", p))
                return True
            except Exception:
                continue

    for base, dirs, files in os.walk("/usr/share/fonts"):
        for f in files:
            if f.endswith(".ttf") and "cjk" in f.lower():
                p = os.path.join(base, f)
                try:
                    pdfmetrics.registerFont(TTFont("CJK-Regular", p))
                    pdfmetrics.registerFont(TTFont("CJK-Bold", p))
                    return True
                except Exception:
                    continue
    return False


def export_plan_pdf(plan_id: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    CJK_OK = _register_cjk_fonts()
    CJK_REGULAR = "CJK-Regular" if CJK_OK else "Helvetica"
    CJK_BOLD = "CJK-Bold" if CJK_OK else "Helvetica-Bold"

    plan = get_plan(plan_id)
    if not plan:
        raise ValueError("Plan not found")

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=20 * mm, bottomMargin=20 * mm)
    styles = getSampleStyleSheet()
    styles["Title"].fontName = CJK_BOLD
    styles["Normal"].fontName = CJK_REGULAR
    elements: list[Any] = []

    # Title
    elements.append(Paragraph(f"志愿填报方案", styles["Title"]))
    elements.append(Spacer(1, 6 * mm))

    # Info section
    info_style = styles["Normal"]
    info_lines = [
        f"考生: {plan.get('user_id', 'default')}",
        f"省份: {plan.get('province', '—')}",
        f"科类: {plan.get('exam_category', '—')}",
        f"位次: {plan.get('rank', '—')}",
        f"状态: {plan.get('status', 'draft')}",
    ]
    for line in info_lines:
        elements.append(Paragraph(line, info_style))
    elements.append(Spacer(1, 4 * mm))

    # Slots table
    slots: list[dict] = plan.get("slots", [])
    if slots:
        header = ["序号", "院校", "专业组", "类别", "录取概率", "组内专业", "说明"]
        data = [header]
        for s in slots:
            prob = _slot_prob(s)
            data.append([
                str(s.get("order", "")),
                s.get("college_name", s.get("college_id", "")),
                s.get("group_code", ""),
                TIER_LABELS.get(s.get("tier", ""), s.get("tier", "")),
                f"{prob * 100:.1f}%",
                _slot_majors_text(s),
                s.get("reason", ""),
            ])

        col_widths = [30, 120, 50, 50, 55, 120, 60]
        table = Table(data, colWidths=col_widths, repeatRows=1)
        table.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, 0), CJK_BOLD),
            ("FONTNAME", (0, 1), (-1, -1), CJK_REGULAR),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F4F6")]),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        elements.append(table)
    else:
        elements.append(Paragraph("志愿表为空", styles["Normal"]))

    doc.build(elements)
    return buf.getvalue()


def export_plan_excel(plan_id: str) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill, Side, Border

    plan = get_plan(plan_id)
    if not plan:
        raise ValueError("Plan not found")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "志愿表"

    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Title row
    ws.cell(row=1, column=1, value="志愿填报方案").font = Font(bold=True, size=14)
    ws.merge_cells("A1:F1")

    # Info rows
    info = [
        ("考生", plan.get("user_id", "default")),
        ("省份", plan.get("province", "—")),
        ("科类", plan.get("exam_category", "—")),
        ("位次", str(plan.get("rank", "—"))),
        ("状态", plan.get("status", "draft")),
    ]
    for i, (k, v) in enumerate(info, 2):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True, size=10)
        ws.cell(row=i, column=2, value=v).font = Font(size=10)

    # Header row
    start_row = len(info) + 3
    headers = ["序号", "院校", "专业组", "类别", "录取概率", "组内专业", "说明"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Data rows
    for i, s in enumerate(plan.get("slots", []), start_row + 1):
        ws.cell(row=i, column=1, value=s.get("order", "")).border = thin_border
        ws.cell(row=i, column=2, value=s.get("college_name", s.get("college_id", ""))).border = thin_border
        ws.cell(row=i, column=3, value=s.get("group_code", "")).border = thin_border
        tier_label = TIER_LABELS.get(s.get("tier", ""), s.get("tier", ""))
        ws.cell(row=i, column=4, value=tier_label).border = thin_border
        ws.cell(row=i, column=5, value=f"{_slot_prob(s) * 100:.1f}%").border = thin_border
        ws.cell(row=i, column=6, value=_slot_majors_text(s)).border = thin_border
        ws.cell(row=i, column=7, value=s.get("reason", "")).border = thin_border

    # Column widths
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 30
    ws.column_dimensions["G"].width = 30

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ============================================================
# 官方志愿样表导出（Phase 23.6）：模板填充 + soffice 转 PDF
# ============================================================

OFFICIAL_TEMPLATE = "data/user/custom/templates/2026_gd_volunteer_form.xlsx"

# 官方样表批次段（行号已实测）：段名 → (起始行, 最大志愿数)
OFFICIAL_SEGMENTS: dict[str, tuple[int, int]] = {
    "提前批本科-空军海军招飞": (6, 1),
    "提前批本科-军检类": (7, 10),
    "提前批本科-艺术类统考+校考": (17, 1),
    "提前批本科-艺术类校考": (17, 1),
    "提前批本科-戏曲类": (18, 1),
    "提前批本科-非军检类": (19, 20),
    "提前批本科-教师专项": (39, 10),
    "提前批本科-卫生专项": (49, 10),
    "艺体类本科批": (81, 20),
    "本科批": (101, 45),
}

# 特殊类型子行（按院校 special_type 定位）
OFFICIAL_SPECIAL_ROWS = {"高水平运动队": 59, "综合评价": 60}


def _major_code(major_id: str) -> str:
    """专业代码：艺体类 major_id 为 {组号}-{代码} 命名空间，去前缀取组内代码。"""
    mid = str(major_id or "")
    return mid.split("-")[-1] if mid else ""


def export_plan_official_excel(plan_id: str) -> bytes:
    """按官方样表模板填入方案 slots → xlsx bytes（布局/合并/样式全保留）。"""
    import openpyxl
    from openpyxl.styles import Alignment, Font
    from pathlib import Path

    plan = get_plan(plan_id)
    if not plan:
        raise ValueError("Plan not found")

    template = Path(__file__).resolve().parents[3] / OFFICIAL_TEMPLATE
    if not template.exists():
        raise FileNotFoundError(f"官方样表模板缺失: {template}")

    wb = openpyxl.load_workbook(str(template))
    ws = wb["Sheet1"]

    batch = plan.get("batch", "本科批")
    slots: list[dict] = plan.get("slots", [])
    slots = sorted(slots, key=lambda s: s.get("order", 0))

    if batch == "提前批本科-特殊类型招生":
        # 特殊类型段：按第一志愿院校的 special_type 定位高水平运动队/综合评价子行
        st = ""
        if slots:
            from deeptutor.services.custom.db import get_connection
            conn = get_connection()
            row = conn.execute(
                "SELECT special_type FROM colleges WHERE id=?", (slots[0].get("college_id", ""),)
            ).fetchone()
            conn.close()
            st = (row["special_type"] if row else "") or ""
        start = OFFICIAL_SPECIAL_ROWS.get(st, OFFICIAL_SPECIAL_ROWS["综合评价"])
        rows = [start]
    else:
        start, cap = OFFICIAL_SEGMENTS.get(batch, OFFICIAL_SEGMENTS["本科批"])
        rows = list(range(start, start + cap))

    for i, slot in enumerate(slots[: len(rows)]):
        r = rows[i]
        ws.cell(r, 6).value = slot.get("province_code") or ""      # F 院校代码（广东招生代码）
        ws.cell(r, 7).value = slot.get("college_name") or ""        # G 院校名称
        gc = str(slot.get("group_code") or "")
        ws.cell(r, 8).value = gc                                    # H 院校专业组代码
        majors = slot.get("majors", [])
        for j, m in enumerate(majors[:6]):
            ws.cell(r, 9 + j).value = _major_code(m.get("major_id"))  # I-N 专业1-6
        ws.cell(r, 15).value = "服从" if slot.get("adjustable", True) else "不服从"  # O
        # 行高按校名长度动态加高 + 院校名称缩小字号：窄列内自动折行显示完整校名（PDF/Excel 同生效）
        name_len = len(str(slot.get("college_name") or ""))
        rows_for_name = max(1, (name_len + 3) // 4)  # G 列宽 ~4字/行（8pt）
        ws.row_dimensions[r].height = 15 + rows_for_name * 14
        ws.cell(r, 7).font = Font(name="Noto Sans CJK SC", size=8)
        ws.cell(r, 7).alignment = Alignment(wrap_text=True, vertical="center")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_plan_official_pdf(plan_id: str) -> bytes:
    """官方样表 xlsx → soffice 转 PDF bytes。"""
    import subprocess
    import tempfile
    from pathlib import Path

    xlsx_bytes = export_plan_official_excel(plan_id)
    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / "form.xlsx"
        src.write_bytes(xlsx_bytes)
        # soffice --headless --convert-to pdf --outdir <dir> <file>
        subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf", "--outdir", td, str(src)],
            check=True, timeout=120,
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
        )
        pdf = Path(td) / "form.pdf"
        if not pdf.exists():
            raise RuntimeError("soffice PDF 转换失败")
        return pdf.read_bytes()
