"""导入 2026 广东提前批本科数据到 admission_ranks。

数据源：专家版 Excel（含提前批本科六类批次）
  - 提前批本科-军检类    561 行
  - 提前批本科-非军检类  249 行
  - 提前批本科-特殊类型招生 188 行
  - 提前批本科-卫生专项  159 行
  - 提前批本科-教师专项  127 行
  - 提前批本科-空军海军招飞  2 行

写入策略：
  - GEN 组级行：year=2026, batch=批次名, min_rank 从 col19 提取 2025 最低位次（无则 0）
  - 专业行：year=2026, batch=批次名, min_rank=0（组级位次兜底）, enrollment_count=2026 计划数
  - 校码经 college_code_map（province='广东'）转官方码

组号：提前批 1xx / 特殊类型 70x，与普通批 2xx 不冲突。

用法：python scripts/import_advance_batch.py
"""
from __future__ import annotations

import re
import sqlite3
import time
from pathlib import Path

import openpyxl

EXCEL_PATH = Path.home() / "桌面/data for agent/广东2026高考志愿大数据专家版0626.xlsx"
ADVANCE_PREFIX = "提前批本科"

GEN_RANK_RE = re.compile(r"最低位次:\s*(\d+)")


def _safe_int(v, default=0) -> int:
    if v is None:
        return default
    try:
        return int(float(str(v).replace(",", "").strip()))
    except (ValueError, TypeError):
        return default


def _safe_str(v, default="") -> str:
    if v is None:
        return default
    return str(v).strip()


def main():
    from deeptutor.services.custom.db import get_connection

    if not EXCEL_PATH.exists():
        print(f"[ERROR] 文件不存在: {EXCEL_PATH}")
        return

    conn = get_connection()
    now = time.time()

    # 校码映射：广东地方码 → 官方码
    code_map = {}
    for r in conn.execute(
        "SELECT province_code, official_code FROM college_code_map WHERE province = ?",
        ("广东",),
    ).fetchall():
        code_map[str(r[0]).strip()] = str(r[1])
    # 官方码自身也认（id 即官方码，含 81/92 段）
    for r in conn.execute("SELECT id FROM colleges").fetchall():
        code_map[str(r[0])] = str(r[0])

    wb = openpyxl.load_workbook(str(EXCEL_PATH), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    group_rows: dict[tuple, dict] = {}  # (official_id, group_code, batch, cat) -> GEN data
    major_rows: list[tuple] = []
    unmatched: dict[str, int] = {}
    total = 0

    for row in ws.iter_rows(min_row=4, values_only=True):
        row = list(row)
        batch = _safe_str(row[3])
        if not batch.startswith(ADVANCE_PREFIX):
            continue
        total += 1
        province_code = _safe_str(row[4])
        exam_category = _safe_str(row[2])
        group_code = _safe_str(row[7])
        major_id = _safe_str(row[9])
        plan = _safe_int(row[16])
        years = _safe_str(row[17])
        tuition = _safe_str(row[18])
        note = _safe_str(row[19])
        subj = _safe_str(row[15])

        official_id = code_map.get(province_code)
        if not official_id:
            unmatched[province_code] = unmatched.get(province_code, 0) + 1
            continue

        # 组级 GEN 行（用 col19 的 2025 最低位次兜底）
        gkey = (official_id, group_code, batch, exam_category)
        if gkey not in group_rows:
            gr = GEN_RANK_RE.search(note)
            min_rank = int(gr.group(1)) if gr else 0
            group_rows[gkey] = {
                "college_id": official_id, "group_code": group_code, "batch": batch,
                "exam_category": exam_category, "min_rank": min_rank,
            }

        # 专业行
        major_rows.append((
            official_id, major_id, "广东", 2026, batch, 0, 0.0, plan,
            exam_category, group_code, years, tuition, subj,
        ))

    wb.close()

    # 写 GEN 组级行
    gen_inserts = 0
    for g in group_rows.values():
        # 用 INSERT OR IGNORE 防止与现有 GEN 行（如普通批同组）冲突
        cur = conn.execute(
            """INSERT OR IGNORE INTO admission_ranks
               (college_id, major_id, province, year, batch, min_rank, min_score,
                enrollment_count, exam_category, group_code)
               VALUES (?, 'GEN', '广东', 2026, ?, ?, 0.0, 0, ?, ?)""",
            (g["college_id"], g["batch"], g["min_rank"], g["exam_category"], g["group_code"]),
        )
        gen_inserts += cur.rowcount

    # 写专业行
    major_inserts = 0
    for m in major_rows:
        official_id, major_id, prov, yr, batch, rk, sc, plan, cat, gcode, yrs, tuit, subj = m
        # 更新 college_major_name 元数据（提前批可能带新专业）
        conn.execute(
            """INSERT OR IGNORE INTO college_major_name
               (college_id, major_id, major_name, full_name, category, subject_requirement,
                tuition, years, campus)
               VALUES (?, ?, '', '', '', ?, ?, ?, '')""",
            (official_id, major_id, subj, _safe_float_0(tuit), yrs),
        )
        cur = conn.execute(
            """INSERT OR IGNORE INTO admission_ranks
               (college_id, major_id, province, year, batch, min_rank, min_score,
                enrollment_count, exam_category, group_code)
               VALUES (?, ?, '广东', 2026, ?, 0, 0.0, ?, ?, ?)""",
            (official_id, major_id, batch, plan, cat, gcode),
        )
        major_inserts += cur.rowcount

    conn.commit()
    conn.close()

    print(f"提前批本科: 处理 {total} 行")
    print(f"  组级 GEN 行: {gen_inserts}（去重后 {len(group_rows)} 组）")
    print(f"  专业行: {major_inserts}")
    if unmatched:
        print(f"  未匹配校码 {len(unmatched)} 个: {sorted(unmatched.items())[:10]}")


def _safe_float_0(v) -> float:
    try:
        return float(str(v).replace(",", "").strip()) if v not in (None, "") else 0.0
    except (ValueError, TypeError):
        return 0.0


if __name__ == "__main__":
    main()