"""Import advisor evaluation data from Excel into the custom database."""

from __future__ import annotations

from pathlib import Path

import openpyxl

from deeptutor.services.custom.advisor_dao import bulk_import
from deeptutor.services.custom.db import init_db

SRC = Path("/home/sunnyxuan2/桌面/data for agent/1.（最全）导师评价表(4月）.xlsx")


def run():
    init_db()
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    evaluations = []
    skip = 0
    for r in rows:
        school, college, name, score, review = (r + (None,) * 7)[:5]
        if not name or not school:
            skip += 1
            continue
        evaluations.append({
            "school": str(school).strip(),
            "college": str(college).strip() if college else None,
            "name": str(name).strip(),
            "score": float(score) if score is not None else 0,
            "review_text": str(review).strip() if review else None,
        })

    wb.close()
    count = bulk_import(evaluations)
    print(f"Imported: {count} advisor evaluations (skipped {skip})")


if __name__ == "__main__":
    run()
